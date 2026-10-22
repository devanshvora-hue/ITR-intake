"""CA workbook download (Stage 5.8 task 3). Synthetic statements, temporary
databases and storage only."""

import io
import logging

import pytest
from openpyxl import load_workbook

import itr_intake_server.app as app_module
import itr_intake_server.case_review as case_review
import itr_intake_server.repository as repo
from itr_intake_server.db import connect
from test_case_review import app, db_path, fast_scrypt, prepare, uploaded, world  # noqa: F401  (fixtures)
from test_client_onboarding import ca_with_code
from test_server_app import db_rows, new_client
from test_statement_uploads import CSV_OK, CSV_PARTLY_BAD, PDFTOTEXT, good_pdf

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
needs_pdf = pytest.mark.skipif(PDFTOTEXT is None, reason="xpdf pdftotext is not available")


def download(client, reference):
    return client.get(f"/api/ca/cases/{reference}/workbook")


def open_book(response):
    return load_workbook(io.BytesIO(response.content))


def summary_values(wb) -> dict:
    ws = wb["Summary"]
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=1).value}


def prepared_case(world, data=CSV_OK, **kw):
    ca, client, reference, _ = world
    statement_id = uploaded(client, reference, data, **kw)
    assert prepare(ca, reference, statement_id).status_code == 200
    return ca, client, reference, statement_id


# === the file =====================================================================================================

def test_a_verified_review_downloads_as_a_valid_attachment(world, db_path):
    ca, _, reference, _ = prepared_case(world)
    response = download(ca, reference)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == f'attachment; filename="{reference}-internal-workbook.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    wb = open_book(response)
    assert wb["Summary"]["A2"].value.startswith("INTERNAL")
    assert wb["Summary"]["A3"].value.startswith("Verified extraction")
    values = summary_values(wb)
    assert values["Client name"] == "Asha Synthetic" and values["Client reference"] == reference
    assert values["Bank"] == "Synthetic Bank Ltd" and values["Account number"] == "XXXX1234"
    assert values["Rows received"] == 2 and values["Rows accepted"] == 2 and values["Rows rejected by validation"] == 0
    assert "s.csv" in values["Data source"] and "Rows not read" not in values["Data source"]
    assert db_rows(db_path, "SELECT count(*) FROM audit_events WHERE action = 'case.workbook_downloaded'")[0][0] == 1


def test_the_financial_year_and_period_match_the_case(world):
    ca, client, reference, _ = prepared_case(world)
    fy = ca.get(f"/api/ca/cases/{reference}").json()["case"]["financial_year"]
    values = summary_values(open_book(download(ca, reference)))
    assert values["Financial year"] == fy and values["Statement period"] == "01-04-2025 to 02-04-2025"


def test_an_unverified_review_is_marked_and_lists_unread_rows(world):
    ca, _, reference, _ = prepared_case(world, CSV_PARTLY_BAD)
    wb = open_book(download(ca, reference))
    banner = wb["Summary"]["A3"].value
    assert banner.startswith("UNVERIFIED REVIEW") and "manual review" in banner and "could not be read" in banner
    assert "Verified extraction" not in banner
    values = summary_values(wb)
    assert values["Rows received"] == 3 and values["Rows accepted"] == 1 and values["Rows rejected by validation"] == 2
    assert "Rows not read: 2 (invalid_date), 3 (debit_invalid_amount)" in values["Data source"]


@needs_pdf
def test_a_verified_pdf_review_downloads(world):
    ca, _, reference, _ = prepared_case(world, good_pdf(), name="s.pdf", ctype="application/pdf")
    wb = open_book(download(ca, reference))
    assert wb["Summary"]["A3"].value.startswith("Verified extraction")
    assert summary_values(wb)["Rows accepted"] == 3


def test_the_workbook_uses_the_upload_the_review_was_built_from(world):
    ca, client, reference, _ = world
    first = uploaded(client, reference, CSV_OK, name="first.csv")
    second = uploaded(client, reference, CSV_PARTLY_BAD, name="second.csv")
    prepare(ca, reference, first)
    prepare(ca, reference, second)
    note = summary_values(open_book(download(ca, reference)))["Data source"]
    assert "second.csv" in note and "first.csv" not in note


def test_the_workbook_is_built_from_stored_rows_not_the_file(world, db_path, tmp_path):
    ca, _, reference, statement_id = prepared_case(world)
    key = db_rows(db_path, "SELECT d.storage_key FROM documents d JOIN statement_uploads u ON u.document_id = d.id "
                           "WHERE u.id = ?", (statement_id,))[0][0]
    (tmp_path / "var" / "storage" / key).unlink()  # temporary test storage only
    response = download(ca, reference)
    assert response.status_code == 200 and summary_values(open_book(response))["Rows accepted"] == 2


FORMULA_CSV = ("Date,Narration,Debit,Credit,Balance\n"
               '01/04/2025,"=HYPERLINK(""http://example.test"",""click"")",100.00,,900.00\n'
               "02/04/2025,+SUM(1+2),50.00,,850.00\n"
               "03/04/2025,@cmd,,25.00,875.00\n"
               "04/04/2025,-2+3 SYNTHETIC,10.00,,865.00\n").encode()


def test_formula_like_descriptions_are_written_as_text_everywhere(world):
    ca, _, reference, _ = prepared_case(world, FORMULA_CSV)
    wb = open_book(download(ca, reference))
    payloads = ("HYPERLINK", "SUM(1+2)", "@cmd", "-2+3 SYNTHETIC")
    hits = [(ws.title, c.coordinate, c.data_type) for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and any(p in c.value for p in payloads)]
    assert hits, "the descriptions should appear in the workbook"
    assert all(dtype == "s" for _, _, dtype in hits), hits


# === access and failures ============================================================================================

def test_no_review_and_unknown_cases_are_not_found(world):
    ca, _, reference, _ = world
    response = download(ca, reference)
    assert response.status_code == 404 and response.json()["error"] == "no_review"
    assert response.headers["cache-control"] == "no-store"
    assert download(ca, "IT-NOPE-000000").status_code == 404
    assert download(ca, "not-a-reference").status_code == 404


def test_anonymous_client_and_other_firm_access_is_refused(app, db_path, world):
    ca, client, reference, _ = prepared_case(world)
    with new_client(app) as anonymous:
        assert download(anonymous, reference).status_code == 401
    assert download(client, reference).status_code == 401  # a client cookie is not a CA session
    other, _ = ca_with_code(app, db_path, "b")
    response = download(other, reference)
    assert response.status_code == 404 and response.headers["content-type"].startswith("application/json")
    other.close()


def test_a_closed_case_can_still_be_downloaded_by_its_firm(world, db_path):
    ca, _, reference, _ = prepared_case(world)
    conn = connect(db_path)
    row = conn.execute("SELECT id, firm_id, version FROM cases").fetchone()
    repo.update_case(conn, repo.FirmScope(row["firm_id"]), row["id"], expected_version=row["version"], status="closed")
    conn.close()
    assert download(ca, reference).status_code == 200


def test_a_generation_failure_is_generic_and_leaks_nothing(app, world, monkeypatch, caplog):
    ca, _, reference, _ = prepared_case(world)

    def boom(*args, **kwargs):
        raise RuntimeError("UPI/1001/SYNTHETIC SHOP 9750.00 C:/secret/path")

    monkeypatch.setattr(case_review, "build_workbook", boom)
    with caplog.at_level(logging.DEBUG):
        response = download(ca, reference)
    assert response.status_code == 500
    assert response.json() == {"error": "workbook_failed", "message": "The workbook could not be generated."}
    for secret in ("SYNTHETIC SHOP", "9750", "secret/path"):
        assert secret not in response.text and secret not in caplog.text


def test_downloads_are_limited_per_firm(world, monkeypatch):
    ca, _, reference, _ = prepared_case(world)
    monkeypatch.setattr(app_module, "WORKBOOKS_PER_FIRM", 1)
    assert download(ca, reference).status_code == 200
    assert download(ca, reference).status_code == 429


def test_the_route_is_get_only(app):
    methods = {m for r in app.routes if getattr(r, "path", "") == "/api/ca/cases/{reference}/workbook"
               for m in r.methods}
    assert methods == {"GET"}
