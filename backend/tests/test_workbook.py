"""Workbook tests. Synthetic data only; files are written to pytest's tmp_path."""

import re
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from itr_intake import (
    CANONICAL_CATEGORIES,
    Category,
    Direction,
    RawTransaction,
    classify_statement,
    validate_statement,
)
from itr_intake.synthetic import BALANCED, UNBALANCED, load_synthetic_statement
from itr_intake.workbook import (
    COL,
    DATE_FORMAT,
    DISCLAIMER,
    INR_FORMAT,
    SHEET_AIS,
    SHEET_ALL,
    SHEET_LISTS,
    SHEET_REVIEW,
    SHEET_SUMMARY,
    STATUS_INCOMPLETE,
    STATUS_NOT_RECONCILED,
    STATUS_RECONCILED,
    StatementMetadata,
    SupportingReportEntry,
    build_workbook,
    mask_account,
    safe_sheet_name,
    save_workbook,
)


def _meta(s):
    return StatementMetadata(
        client_name=s.profile.name,
        client_reference="DEMO-0142",
        bank_name="Example Bank",
        account_number="50100012345678",
        financial_year="FY 2025-26",
        period_start=s.period_start,
        period_end=s.period_end,
        data_source_note="SYNTHETIC DEMO DATA",
    )


def generate(tmp_path, filename=BALANCED, supporting=None):
    s = load_synthetic_statement(filename)
    report = validate_statement(s.rows, s.opening_balance, s.closing_balance,
                                period_start=s.period_start, period_end=s.period_end)
    classified = classify_statement(report.transactions, s.profile)
    path = save_workbook(
        build_workbook(report, classified, _meta(s), supporting_reports=supporting),
        tmp_path / f"{filename}.xlsx",
    )
    return load_workbook(path), report, classified


@pytest.fixture
def balanced(tmp_path):
    return generate(tmp_path)


def find_row(ws, label):
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if cell.value == label:
            return cell.row
    raise AssertionError(f"label not found: {label}")


# A tiny evaluator for the exact formula shapes the generator writes, so the
# tests check that the live formulas point at the right data.
_RANGE = r"'All Transactions'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)"


def _column(ws, letter, first, last):
    from openpyxl.utils import column_index_from_string
    c = column_index_from_string(letter)
    return [ws.cell(row=r, column=c).value for r in range(int(first), int(last) + 1)]


def evaluate(wb, formula, row_label=None):
    all_ws = wb[SHEET_ALL]
    if m := re.fullmatch(rf"=COUNTIF\({_RANGE},A\d+\)", formula):
        cats = _column(all_ws, m[1], m[2], m[4])
        return sum(1 for c in cats if c == row_label)
    if m := re.fullmatch(rf"=SUMIFS\({_RANGE},{_RANGE},(?:A\d+|\"([^\"]+)\")\)", formula):
        values = _column(all_ws, m[1], m[2], m[4])
        cats = _column(all_ws, m[5], m[6], m[8])
        target = m[9] if m[9] is not None else row_label
        return sum((Decimal(str(v)) for v, c in zip(values, cats) if c == target and v is not None), Decimal("0"))
    if m := re.fullmatch(rf"=SUM\({_RANGE}\)", formula):
        return sum((Decimal(str(v)) for v in _column(all_ws, m[1], m[2], m[4]) if v is not None), Decimal("0"))
    raise AssertionError(f"unsupported formula: {formula}")


def totals_by_category(classified):
    out = {}
    for t in classified:
        n, d, c = out.get(t.category.value, (0, Decimal("0"), Decimal("0")))
        out[t.category.value] = (n + 1, d + t.debit, c + t.credit)
    return out


# --- sheet names ---------------------------------------------------------------

def test_sheet_names_and_order(balanced):
    wb, _, classified = balanced
    used = [c for c in CANONICAL_CATEGORIES if any(t.category.value == c for t in classified)]
    expected_category_sheets = [safe_sheet_name(c, set()) for c in used]
    assert wb.sheetnames == [SHEET_SUMMARY, SHEET_ALL, *expected_category_sheets, SHEET_REVIEW, SHEET_LISTS]
    assert wb[SHEET_LISTS].sheet_state == "hidden"
    assert SHEET_AIS not in wb.sheetnames  # no supporting data supplied
    for name in wb.sheetnames:
        assert len(name) <= 31 and not set(name) & set("[]:*?/\\")


def test_zero_categories_get_no_sheet(balanced):
    wb, _, _ = balanced
    assert "Dividend" not in wb.sheetnames
    assert "Rent Received" not in wb.sheetnames


def test_safe_sheet_name_rules():
    assert safe_sheet_name("Salary / Pension", set()) == "Salary - Pension"
    long = safe_sheet_name("Loans Received / Repayment of Loans Given", set())
    assert len(long) <= 31 and "/" not in long
    assert safe_sheet_name("Summary", set()) == "Summary (2)"
    assert safe_sheet_name("needs review", set()) == "needs review (2)"
    assert safe_sheet_name("A", {"a"}) == "A (2)"
    assert safe_sheet_name("'quoted'", set()) == "quoted"


# --- All Transactions: preservation and formats ---------------------------------

def test_all_transactions_preserve_source_values_in_date_order(balanced):
    wb, report, classified = balanced
    ws = wb[SHEET_ALL]
    assert [c.value for c in ws[1]][:4] == ["Row #", "Date", "Mode", "Counterparty"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == report.valid_count == 24
    dates = [r[COL["Date"] - 1] for r in rows]
    assert dates == sorted(dates)
    by_row = {t.row_number: t for t in classified}
    for r in rows:
        t = by_row[r[COL["Row #"] - 1]]
        assert r[COL["Narration"] - 1] == t.narration
        assert r[COL["Date"] - 1].date() == t.date
        debit, credit = r[COL["Debit (₹)"] - 1], r[COL["Credit (₹)"] - 1]
        assert (Decimal(str(debit)) if debit is not None else Decimal("0")) == t.debit
        assert (Decimal(str(credit)) if credit is not None else Decimal("0")) == t.credit
        assert Decimal(str(r[COL["Balance (₹)"] - 1])) == t.balance
        assert r[COL["Category"] - 1] == t.category.value
        assert r[COL["Classification reason"] - 1] == t.classification_reason
        assert r[COL["Review status"] - 1] == t.review_status.value


def test_number_and_date_formats(balanced):
    wb, _, _ = balanced
    ws = wb[SHEET_ALL]
    assert ws.cell(row=2, column=COL["Date"]).number_format == DATE_FORMAT == "DD-MM-YYYY"
    for name in ("Debit (₹)", "Credit (₹)", "Balance (₹)"):
        assert ws.cell(row=2, column=COL[name]).number_format == INR_FORMAT
    assert "##\\,##\\,##0.00" in INR_FORMAT  # lakh grouping
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref.startswith("A1:")
    assert ws.cell(row=1, column=1).font.bold

    cat_ws = wb["Salary - Pension"]
    assert [c.value for c in cat_ws[2]] == [
        "Date", "Mode", "Counterparty", "Debit (₹)", "Credit (₹)", "Balance (₹)", "Narration", "Remarks",
    ]
    assert cat_ws.cell(row=3, column=1).number_format == DATE_FORMAT
    assert cat_ws.cell(row=3, column=5).number_format == INR_FORMAT
    assert cat_ws.freeze_panes == "A3"


def test_category_column_has_dropdown_of_28_categories(balanced):
    wb, _, _ = balanced
    dvs = wb[SHEET_ALL].data_validations.dataValidation
    assert len(dvs) == 1 and dvs[0].type == "list"
    assert "Lists" in dvs[0].formula1
    assert [wb[SHEET_LISTS].cell(row=i, column=1).value for i in range(1, 29)] == list(CANONICAL_CATEGORIES)


def test_formula_like_narration_is_stored_as_text(tmp_path):
    rows = [RawTransaction(1, "01/04/2025", "=HYPERLINK(\"x\")", "", "100.00")]
    report = validate_statement(rows, "0", "100.00")
    from itr_intake import ClientProfile
    classified = classify_statement(report.transactions, ClientProfile(name="Demo Person"))
    path = save_workbook(build_workbook(report, classified, StatementMetadata("Demo Person", "DEMO-1")), tmp_path / "x.xlsx")
    cell = load_workbook(path)[SHEET_ALL].cell(row=2, column=COL["Narration"])
    assert cell.value == "=HYPERLINK(\"x\")" and cell.data_type == "s"


# --- Summary: totals and formulas -------------------------------------------------

def test_summary_category_totals_are_live_and_correct(balanced):
    wb, _, classified = balanced
    ws = wb[SHEET_SUMMARY]
    expected = totals_by_category(classified)
    for name in CANONICAL_CATEGORIES:
        r = find_row(ws, name)
        count_f, debit_f, credit_f, net_f = (ws.cell(row=r, column=c).value for c in range(2, 6))
        assert count_f.startswith("=COUNTIF(") and debit_f.startswith("=SUMIFS(")
        assert net_f == f"=D{r}-C{r}"
        n, d, c = expected.get(name, (0, Decimal("0"), Decimal("0")))
        assert evaluate(wb, count_f, name) == n
        assert evaluate(wb, debit_f, name) == d
        assert evaluate(wb, credit_f, name) == c
        assert ws.cell(row=r, column=3).number_format == INR_FORMAT

    total = find_row(ws, "Total")
    first, last = find_row(ws, CANONICAL_CATEGORIES[0]), find_row(ws, CANONICAL_CATEGORIES[-1])
    assert last - first == 27
    assert ws.cell(row=total, column=2).value == f"=SUM(B{first}:B{last})"
    assert ws.cell(row=total, column=4).value == f"=SUM(D{first}:D{last})"
    assert ws.cell(row=total, column=2).font.bold


def test_category_sheet_rows_and_totals(balanced):
    wb, _, classified = balanced
    ws = wb["Salary - Pension"]
    salary = [t for t in classified if t.category is Category.SALARY_PENSION]
    data = list(ws.iter_rows(min_row=3, max_row=2 + len(salary), values_only=True))
    assert [Decimal(str(r[4])) for r in data] == [t.credit for t in salary]
    total_row = 3 + len(salary)
    assert ws.cell(row=total_row, column=1).value == "Total"
    assert ws.cell(row=total_row, column=5).value == f"=SUM(E3:E{total_row - 1})"


def test_key_figures_are_formulas(balanced):
    wb, _, classified = balanced
    ws = wb[SHEET_SUMMARY]
    f = ws.cell(row=find_row(ws, "Credits classified as Salary / Pension"), column=2).value
    assert evaluate(wb, f) == sum(t.credit for t in classified if t.category is Category.SALARY_PENSION)


def test_summary_metadata_and_disclaimer(balanced):
    wb, report, _ = balanced
    ws = wb[SHEET_SUMMARY]
    def value(label):
        return ws.cell(row=find_row(ws, label), column=2).value
    assert value("Client reference") == "DEMO-0142"
    assert value("Bank") == "Example Bank"
    assert value("Account number") == "XXXX5678"
    assert value("Financial year") == "FY 2025-26"
    assert value("Statement period") == "01-04-2025 to 30-06-2025"
    assert value("Rows received") == report.rows_received
    assert value("Rows rejected by validation") == 0
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert DISCLAIMER in texts
    assert any("INTERNAL" in t and "client" in t for t in texts)


def test_mask_account():
    assert mask_account("50100012345678") == "XXXX5678"
    assert mask_account(None) == "Not provided"


# --- reconciliation ---------------------------------------------------------------

def test_successful_reconciliation(balanced):
    wb, report, _ = balanced
    ws = wb[SHEET_SUMMARY]
    def cell(label):
        return ws.cell(row=find_row(ws, label), column=2)
    assert cell("Status at generation").value == STATUS_RECONCILED
    assert Decimal(str(cell("Difference at generation").value)) == Decimal("0")
    assert evaluate(wb, cell("Total credits").value) == report.reconciliation.total_credits
    assert evaluate(wb, cell("Total debits").value) == report.reconciliation.total_debits
    opening, credits, debits = (cell(l).coordinate for l in ("Opening balance (statement)", "Total credits", "Total debits"))
    assert cell("Expected closing balance").value == f"={opening}+{credits}-{debits}"
    assert cell("Difference (statement − expected)").value.startswith("=")
    live = cell("Status (live)").value
    assert STATUS_RECONCILED in live and STATUS_NOT_RECONCILED in live and "ABS(" in live
    assert cell("Opening balance (statement)").number_format == INR_FORMAT


def test_failed_reconciliation_is_prominent_and_unadjusted(tmp_path):
    wb, report, _ = generate(tmp_path, UNBALANCED)
    ws = wb[SHEET_SUMMARY]
    status = ws.cell(row=find_row(ws, "Status at generation"), column=2)
    diff = ws.cell(row=find_row(ws, "Difference at generation"), column=2)
    assert status.value == STATUS_NOT_RECONCILED
    assert status.fill.fgColor.rgb.endswith("F5D0C8") and status.font.bold
    assert Decimal(str(diff.value)) == Decimal("1000.00") == report.reconciliation.difference
    stated = ws.cell(row=find_row(ws, "Closing balance (statement)"), column=2).value
    assert Decimal(str(stated)) == report.reconciliation.stated_closing_balance
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("does NOT reconcile" in t and "1,000.00" in t and "No amounts have been adjusted" in t for t in texts)
    assert STATUS_RECONCILED != status.value


def test_rejected_rows_show_incomplete_status(tmp_path):
    rows = [
        RawTransaction(1, "01/04/2025", "SYNTHETIC CREDIT", "", "100.00"),
        RawTransaction(2, "31/02/2025", "SYNTHETIC BAD DATE", "", "5.00"),
    ]
    report = validate_statement(rows, "0", "100.00")
    from itr_intake import ClientProfile
    classified = classify_statement(report.transactions, ClientProfile(name="Demo Person"))
    path = save_workbook(build_workbook(report, classified, StatementMetadata("Demo Person", "DEMO-2")), tmp_path / "i.xlsx")
    ws = load_workbook(path)[SHEET_SUMMARY]
    assert ws.cell(row=find_row(ws, "Status at generation"), column=2).value == STATUS_INCOMPLETE
    assert ws.cell(row=find_row(ws, "Rows rejected by validation"), column=2).value == 1


# --- Needs Review -----------------------------------------------------------------

def test_needs_review_contents(balanced):
    wb, _, classified = balanced
    ws = wb[SHEET_REVIEW]
    header = [c.value for c in ws[1]]
    assert "Suggested category" in header and "Reason" in header
    listed = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    expected = {
        t.row_number for t in classified
        if t.review_status.value != "Auto-classified" or t.flags or t.category is Category.OTHERS_NEEDS_REVIEW
    }
    assert set(listed) == expected
    assert {7, 8, 9, 19, 24} <= set(listed)
    assert 1 not in listed  # salary from a listed employer is auto-classified

    large = listed[7]
    assert "large_credit" in large[header.index("Flags")]
    assert large[header.index("Suggested category")].startswith("Not determined")
    assert large[header.index("Reason")]
    family = listed[8]
    assert family[header.index("Suggested category")].startswith("Family Transfer")


def test_needs_review_suggestions_are_not_tax_conclusions(balanced):
    wb, _, _ = balanced
    for row in wb[SHEET_REVIEW].iter_rows(min_row=2, values_only=True):
        text = " ".join(str(v) for v in row if v).lower()
        assert "taxable" not in text and "deductible" not in text


# --- AIS / 26AS ---------------------------------------------------------------------

def test_ais_sheet_only_when_data_supplied(tmp_path):
    entries = [
        SupportingReportEntry("AIS", "Savings bank interest (synthetic)", Decimal("1245.00"),
                              Category.INTEREST_SAVINGS, Direction.CREDIT),
        SupportingReportEntry("26AS", "TDS on salary (synthetic)", Decimal("5000.00")),
    ]
    wb, _, _ = generate(tmp_path, supporting=entries)
    assert SHEET_AIS in wb.sheetnames
    ws = wb[SHEET_AIS]
    assert ws.cell(row=2, column=5).value.startswith("=SUMIFS(")
    assert evaluate(wb, ws.cell(row=2, column=5).value.replace("D2", '"Interest – Savings Bank"')) == Decimal("1245.00")
    assert ws.cell(row=2, column=6).value == "=C2-E2"
    assert re.search(r"'?Summary'?!\$?B\$?\d+", ws.cell(row=2, column=7).value)
    assert ws.cell(row=3, column=4).value == "Not mapped"
    assert ws.cell(row=3, column=7).value == "Not compared"


# --- safety -----------------------------------------------------------------------

def test_save_refuses_private_folder(tmp_path, balanced):
    wb, _, _ = balanced
    with pytest.raises(ValueError):
        save_workbook(wb, tmp_path / "Private" / "x.xlsx")


def test_build_requires_classified_transactions():
    s = load_synthetic_statement()
    report = validate_statement(s.rows, s.opening_balance, s.closing_balance)
    with pytest.raises(ValueError):
        build_workbook(report, list(report.transactions), _meta(s))


def test_build_does_not_modify_classified_transactions(tmp_path):
    s = load_synthetic_statement()
    report = validate_statement(s.rows, s.opening_balance, s.closing_balance)
    classified = classify_statement(report.transactions, s.profile)
    snapshot = list(classified)
    build_workbook(report, classified, _meta(s))
    assert classified == snapshot
    assert date(2025, 4, 1) == classified[0].date
